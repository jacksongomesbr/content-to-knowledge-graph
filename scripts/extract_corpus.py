"""Extract document text to JSONL units with locators. No semantic graph extraction."""
from cli_common import Parser, emit, run
import argparse
import hashlib
import json
from pathlib import Path

SUPPORTED = {'.txt','.md','.qmd','.html','.htm','.pdf','.docx','.epub'}
SKIP = {'.generated', '.wikipedia-cache', '.git','node_modules','.venv','dist','_output','__pycache__'}


def iter_extract(path):
    ext = path.suffix.lower()
    if ext == '.pdf':
        from pypdf import PdfReader
        for i,page in enumerate(PdfReader(path).pages): yield i+1,page.extract_text() or ''
        return
    if ext == '.docx':
        from docx import Document
        doc = Document(path)
        yield 'document', '\n'.join([p.text for p in doc.paragraphs] + [' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        return
    if ext == '.epub':
        import zipfile
        from bs4 import BeautifulSoup
        import xml.etree.ElementTree as ET
        with zipfile.ZipFile(path) as z:
            container = ET.fromstring(z.read('META-INF/container.xml'))
            opf = next(e.attrib['full-path'] for e in container.iter() if e.tag.endswith('rootfile'))
            root = ET.fromstring(z.read(opf)); manifest = {e.attrib['id']:e.attrib['href'] for e in root.iter() if e.tag.endswith('item')}
            import posixpath
            for e in root.iter():
                if e.tag.endswith('itemref'):
                    name = posixpath.normpath(posixpath.join(posixpath.dirname(opf), manifest[e.attrib['idref']]))
                    soup=BeautifulSoup(z.read(name),'html.parser')
                    for element in soup(['script','style']): element.decompose()
                    yield name,soup.get_text('\n')
            return
    text = path.read_text(encoding='utf-8-sig')
    if ext in ('.html','.htm'):
        from bs4 import BeautifulSoup
        soup = BeautifulSoup(text, 'html.parser')
        for el in soup(['script','style']): el.decompose()
        text = soup.get_text('\n')
    yield 'document', text


def extract(path):
    # Compatibility API; command-line extraction streams units instead.
    return list(iter_extract(path))


def main():
    p=Parser(description=__doc__); p.add_argument('origins', nargs='+', type=Path); p.add_argument('--output', required=True, type=Path)
    a=p.parse_args(); output=a.output.resolve()
    if output.exists() and any(output.iterdir()): p.error('Use uma pasta de extração nova/vazia para preservar resultados existentes.')
    output.mkdir(parents=True,exist_ok=True); coverage=[]; seen=set(); unit_count=0
    from evidence_passages import file_sha256
    with (output/'units.jsonl').open('x',encoding='utf-8') as unit_stream:
        for origin in a.origins:
            origin=origin.resolve()
            if not origin.exists(): p.error(f'Origem inexistente: {origin}')
            files=sorted(origin.rglob('*')) if origin.is_dir() else [origin]
            matched=0
            for file in files:
                if not file.is_file() or file.is_symlink() or file.suffix.lower() not in SUPPORTED or any(x in SKIP for x in file.parts): continue
                actual=file.resolve()
                if output == actual or output in actual.parents: continue
                if actual in seen: matched+=1; continue
                if origin.is_dir() and not actual.is_relative_to(origin): continue
                seen.add(actual); matched+=1
                item={'path':str(actual)};nonempty=total=0
                try:
                    before=file.stat(); checksum=file_sha256(file)
                    item['sha256']=checksum
                    for locator,text in iter_extract(file):
                        total+=1
                        if not text.strip(): continue
                        nonempty+=1
                        uid=hashlib.sha256(f'{actual}:{locator}'.encode()).hexdigest()[:20]
                        unit=dict(id=uid,sourceFile=str(actual),locator=locator,text=text,textSha256=hashlib.sha256(text.encode()).hexdigest())
                        unit_stream.write(json.dumps(unit,ensure_ascii=False)+'\n');unit_count+=1
                    item.update(status='read' if nonempty == total and total else 'partial', units=nonempty, totalUnits=total)
                    after=file.stat()
                    if (before.st_mtime_ns,before.st_ctime_ns,before.st_size)!=(after.st_mtime_ns,after.st_ctime_ns,after.st_size):
                        item.update(status='partial',limitation='A fonte mudou durante a leitura; descarte as unidades deste arquivo e repita a extração.')
                    elif nonempty != total or not total: item['limitation']='Unidades sem texto extraível; OCR pode ser necessário.'
                except Exception as exc: item.update(status='partial' if nonempty else 'blocked',error=str(exc),units=nonempty,totalUnits=total)
                coverage.append(item)
            if not matched and not any(str(origin)==str(entry.get('path')) for entry in coverage):
                coverage.append(dict(path=str(origin),status='blocked',error='Nenhum arquivo novo em formato suportado foi encontrado.'))
    (output/'coverage.json').write_text(json.dumps(coverage,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    partial = not unit_count or any(c['status']!='read' for c in coverage)
    emit(dict(output=str(output), units=unit_count, files=len(coverage)), human=f'{unit_count} unidades; {len(coverage)} arquivos. Extração textual não equivale a revisão semântica.', status='partial' if partial else 'success')
    return 3 if partial else 0

if __name__ == '__main__': raise SystemExit(run(main))
